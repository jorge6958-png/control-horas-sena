import os
import json
import re
import pandas as pd
import streamlit as st
from rapidfuzz import fuzz

RUTA_JSON = os.path.join(os.path.dirname(__file__), "competencias.json")


def cargar_competencias(programa):
    with open(RUTA_JSON, encoding="utf-8") as f:
        data = json.load(f)
    prog = data["programas"].get(programa)
    if not prog:
        raise ValueError(f"Programa no encontrado: {programa}")
    df = pd.DataFrame(prog["competencias"])
    return df


def listar_programas():
    with open(RUTA_JSON, encoding="utf-8") as f:
        data = json.load(f)
    return list(data["programas"].keys())


def match_competencia_fuzzy(texto, lista_ref, threshold=82):
    texto_clean = " ".join(texto.lower().split())
    best_match = None
    best_score = 0
    for ref in lista_ref:
        ref_clean = " ".join(ref.lower().split())
        score = fuzz.partial_ratio(ref_clean, texto_clean)
        if score > best_score:
            best_score = score
            best_match = ref
    return best_match if best_score >= threshold else None


def asignar_competencias(df_detalle, df_competencias):
    df_detalle = df_detalle.copy()
    ref_nombres = list(df_competencias["nombre"])
    ref_set = set(n.upper().strip() for n in ref_nombres)

    df_detalle["competencia_normalizada"] = ""
    df_detalle["coincide"] = False

    for idx, row in df_detalle.iterrows():
        xls = row["competencia_xls"]
        xls_lower = " ".join(xls.lower().split())

        # 1. Exact match
        for ref_name in ref_nombres:
            if xls_lower == " ".join(ref_name.lower().split()):
                df_detalle.at[idx, "competencia_normalizada"] = ref_name
                df_detalle.at[idx, "coincide"] = True
                break
        if df_detalle.at[idx, "coincide"]:
            continue

        # 2. Substring match
        for ref_name in ref_nombres:
            ref_lower = " ".join(ref_name.lower().split())
            if ref_lower in xls_lower or xls_lower in ref_lower:
                df_detalle.at[idx, "competencia_normalizada"] = ref_name
                df_detalle.at[idx, "coincide"] = True
                break
        if df_detalle.at[idx, "coincide"]:
            continue

        # 3. Fuzzy match
        fuzzy_match = match_competencia_fuzzy(xls, ref_nombres)
        if fuzzy_match:
            df_detalle.at[idx, "competencia_normalizada"] = fuzzy_match
            df_detalle.at[idx, "coincide"] = True

    return df_detalle


def procesar_reporte(ruta_archivo):
    ext = os.path.splitext(ruta_archivo)[1].lower()

    if ext in (".xls",):
        import xlrd
        wb = xlrd.open_workbook(ruta_archivo)
        sheet = wb.sheet_by_index(0)

        info_ficha = {}
        for i in range(1, min(8, sheet.nrows)):
            label = str(sheet.cell_value(i, 0)).strip()
            raw = sheet.cell_value(i, 1) if sheet.ncols > 1 else ""
            if isinstance(raw, float) and raw == int(raw):
                raw = int(raw)
            valor = str(raw).strip()
            if label and valor:
                info_ficha[label] = valor

        filas = []
        for i in range(11, sheet.nrows):
            nombre = str(sheet.cell_value(i, 0)).strip()
            apellido = str(sheet.cell_value(i, 1)).strip()
            estado = str(sheet.cell_value(i, 2)).strip()
            competencia = str(sheet.cell_value(i, 3)).strip()
            fecha_ini = str(sheet.cell_value(i, 4)).strip()
            fecha_fin = str(sheet.cell_value(i, 5)).strip()
            horas_str = str(sheet.cell_value(i, 6)).strip()

            if not nombre or not competencia:
                continue

            try:
                horas = round(float(horas_str)) if horas_str else 0
            except ValueError:
                horas = 0

            filas.append({
                "nombre": nombre,
                "apellido": apellido,
                "instructor": f"{nombre} {apellido}",
                "estado": estado,
                "competencia_xls": competencia,
                "fecha_inicio": fecha_ini,
                "fecha_fin": fecha_fin,
                "horas": horas,
            })

    elif ext in (".xlsx",):
        import openpyxl
        wb = openpyxl.load_workbook(ruta_archivo, read_only=True, data_only=True)
        ws = wb.active

        info_ficha = {}
        for i in range(1, min(8, ws.max_row + 1)):
            c0 = str(ws.cell(i, 1).value or "").strip()
            raw = ws.cell(i, 2).value
            if isinstance(raw, float) and raw == int(raw):
                raw = int(raw)
            c1 = str(raw or "").strip()
            if c0 and c1:
                info_ficha[c0] = c1

        filas = []
        for i in range(11, ws.max_row + 1):
            nombre = str(ws.cell(i, 1).value or "").strip()
            apellido = str(ws.cell(i, 2).value or "").strip()
            estado = str(ws.cell(i, 3).value or "").strip()
            competencia = str(ws.cell(i, 4).value or "").strip()
            fecha_ini = str(ws.cell(i, 5).value or "").strip()
            fecha_fin = str(ws.cell(i, 6).value or "").strip()
            horas_str = str(ws.cell(i, 7).value or "").strip()

            if not nombre or not competencia:
                continue

            try:
                horas = round(float(horas_str)) if horas_str else 0
            except ValueError:
                horas = 0

            filas.append({
                "nombre": nombre,
                "apellido": apellido,
                "instructor": f"{nombre} {apellido}",
                "estado": estado,
                "competencia_xls": competencia,
                "fecha_inicio": fecha_ini,
                "fecha_fin": fecha_fin,
                "horas": horas,
            })

    elif ext in (".csv",):
        df_csv = pd.read_csv(ruta_archivo, encoding="utf-8", sep=None, engine="python")
        info_ficha = {}
        filas = []

        for i, row in df_csv.iterrows():
            valores = [str(v).strip() for v in row]
            if i < 8 and len(valores) >= 2 and valores[0]:
                info_ficha[valores[0]] = valores[1]
                continue
            if i < 10:
                continue

            if len(valores) >= 7:
                nombre = valores[0]
                apellido = valores[1]
                estado = valores[2]
                competencia = valores[3]
                fecha_ini = valores[4]
                fecha_fin = valores[5]
                horas_str = valores[6]
            else:
                continue

            if not nombre or not competencia:
                continue

            try:
                horas = round(float(horas_str)) if horas_str else 0
            except ValueError:
                horas = 0

            filas.append({
                "nombre": nombre,
                "apellido": apellido,
                "instructor": f"{nombre} {apellido}",
                "estado": estado,
                "competencia_xls": competencia,
                "fecha_inicio": fecha_ini,
                "fecha_fin": fecha_fin,
                "horas": horas,
            })
    else:
        raise ValueError(f"Formato no soportado: {ext}. Usa .xls, .xlsx o .csv")

    df_detalle = pd.DataFrame(filas)
    return df_detalle, info_ficha


def construir_tabla_competencias(df_detalle, df_competencias):
    df_ref = df_competencias.copy()

    if not df_detalle.empty and "competencia_normalizada" in df_detalle.columns:
        coincidentes = df_detalle[df_detalle["coincide"]]

        df_rep = coincidentes.groupby("competencia_normalizada").agg(
            horas_reportadas=("horas", "sum"),
        ).reset_index()
        if not df_rep.empty:
            df_rep.columns = ["nombre", "horas_reportadas"]
        else:
            df_rep = pd.DataFrame(columns=["nombre", "horas_reportadas"])

        instr_estados = coincidentes.groupby("competencia_normalizada").apply(
            lambda g: "; ".join(
                f"{row['instructor']} ({row['estado']})"
                for _, row in g.iterrows()
            ),
            include_groups=False,
        ).reset_index()
        if not instr_estados.empty:
            instr_estados.columns = ["nombre", "instructores_detalle"]
        else:
            instr_estados = pd.DataFrame(columns=["nombre", "instructores_detalle"])
    else:
        df_rep = pd.DataFrame(columns=["nombre", "horas_reportadas"])
        instr_estados = pd.DataFrame(columns=["nombre", "instructores_detalle"])

    df_ref["key"] = df_ref["nombre"].str.upper().str.strip()
    df_rep["key"] = df_rep["nombre"].str.upper().str.strip()
    instr_estados["key"] = instr_estados["nombre"].str.upper().str.strip()

    df_result = df_ref.merge(
        df_rep[["key", "horas_reportadas"]],
        on="key",
        how="left",
    )
    df_result = df_result.merge(
        instr_estados[["key", "instructores_detalle"]],
        on="key",
        how="left",
    )

    df_result["horas_reportadas"] = df_result["horas_reportadas"].fillna(0).astype(int)
    df_result["diferencia"] = df_result["horas_reportadas"] - df_result["horas_planeadas"]
    df_result["porcentaje"] = (
        (df_result["horas_reportadas"] / df_result["horas_planeadas"] * 100).round(1)
    )
    df_result["instructores_detalle"] = df_result["instructores_detalle"].fillna("\u2014")

    df_result = df_result.sort_values("horas_planeadas", ascending=False).reset_index(drop=True)
    return df_result


def construir_tabla_instructores(df_detalle):
    if df_detalle.empty:
        return pd.DataFrame()

    competencias_map = {}
    for _, row in df_detalle.iterrows():
        instr = row["instructor"]
        if instr not in competencias_map:
            competencias_map[instr] = set()
        competencias_map[instr].add(row["competencia_xls"])

    df = df_detalle.groupby(["instructor", "estado"]).agg(
        total_horas=("horas", "sum"),
        registros=("horas", "count"),
    ).reset_index()

    df["competencias"] = df["instructor"].map(
        lambda x: "; ".join(sorted(competencias_map.get(x, set())))
    )

    df = df.sort_values("total_horas", ascending=False).reset_index(drop=True)
    return df
